"""Tests for multi-format resume import, discovery, validation, and claims."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from fpdf import FPDF

from src.common.profile_review import review_profile
from src.common.profile_validation import validate_profile
from src.common.resume_discovery import discover_resume, list_resume_files
from src.common.resume_import import (
    extract_metric,
    import_resume,
    parse_dates,
    split_sections,
)
from src.common.resume_readers import extract_resume_text, normalize_resume_text
from src.job_parser.parser import parse_job_dict
from src.matching.scorer import MatchScorer
from src.models.candidate import (
    CandidateProfile,
    ClaimItem,
    ClaimType,
    EducationRecord,
    SkillRecord,
    VerificationStatus,
    VerifiedExperience,
)
from src.resume_tailoring.tailor import ResumeTailor


FIXTURES = Path(__file__).parent / "fixtures"


def _write_pdf(path: Path, text: str) -> None:
    pdf = FPDF(format="Letter")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_margins(20, 20, 20)
    pdf.set_font("Helvetica", size=10)
    for line in text.splitlines() or [""]:
        safe = (
            line.replace("•", "-")
            .replace("–", "-")
            .replace("—", "-")
            .replace("’", "'")
            .replace("‘", "'")
            .replace("“", '"')
            .replace("”", '"')
        )
        # Helvetica is Latin-1; drop any remaining unsupported glyphs
        safe = safe.encode("latin-1", "replace").decode("latin-1")
        pdf.cell(0, 6, safe, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(path))


def _write_docx(path: Path, text: str) -> None:
    doc = Document()
    for line in text.splitlines():
        doc.add_paragraph(line)
    doc.save(str(path))


@pytest.fixture
def sample_text() -> str:
    return (FIXTURES / "sample_resume.txt").read_text(encoding="utf-8")


def test_markdown_import(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    from src.common import resume_import as ri

    monkeypatch.setattr(ri, "CANDIDATE_DIR", tmp_path / "candidate")
    monkeypatch.setattr(ri, "CONFIG_DIR", tmp_path / "config")
    monkeypatch.setattr(ri, "RESUMES_MASTER", tmp_path / "resumes")
    (tmp_path / "candidate").mkdir()
    (tmp_path / "config").mkdir()
    (tmp_path / "resumes").mkdir()
    # minimal profile
    (tmp_path / "config" / "candidate_profile.yaml").write_text(
        "full_name: Jaideep Ponnam\npreferred_name: Jai\nlocation:\n  city: Austin\n  state: Texas\n",
        encoding="utf-8",
    )

    result = import_resume(FIXTURES / "sample_resume.md", dry_run=False)
    assert result["experiences_found"] >= 1
    assert result["skills_found"] >= 5
    assert result["format"] == ".md"
    # unsupported claim detected
    exp_yaml = (tmp_path / "candidate" / "verified_experience.yaml").read_text(encoding="utf-8")
    assert "unsupported" in exp_yaml
    assert "35%" in exp_yaml  # metric preserved


def test_txt_import_dry_run():
    result = import_resume(FIXTURES / "sample_resume.txt", dry_run=True)
    assert result["dry_run"] is True
    assert result["experiences_found"] >= 1
    assert result["skills_found"] >= 3
    assert result["education_blocks_found"] >= 1
    assert result["certifications_found"] >= 1
    preview = result["preview"]["experiences"][0]
    assert preview["start_date"] == "Jun 2023"
    assert preview["end_date"] == "Aug 2023"
    assert any("35%" in (m.get("text") + str(m.get("metric"))) for m in preview["measurable_results"])


def test_pdf_import(tmp_path: Path, sample_text: str, monkeypatch: pytest.MonkeyPatch):
    from src.common import resume_import as ri

    monkeypatch.setattr(ri, "CANDIDATE_DIR", tmp_path / "candidate")
    monkeypatch.setattr(ri, "CONFIG_DIR", tmp_path / "config")
    monkeypatch.setattr(ri, "RESUMES_MASTER", tmp_path / "resumes")
    (tmp_path / "candidate").mkdir()
    (tmp_path / "config").mkdir()
    (tmp_path / "resumes").mkdir()
    (tmp_path / "config" / "candidate_profile.yaml").write_text(
        "full_name: Jaideep Ponnam\nlocation:\n  city: Austin\n  state: Texas\n",
        encoding="utf-8",
    )
    pdf_path = tmp_path / "resume.pdf"
    _write_pdf(pdf_path, sample_text)
    text = extract_resume_text(pdf_path)
    assert "Jaideep" in text or "JAIDEEP" in text.upper()
    result = import_resume(pdf_path, dry_run=True)
    assert result["format"] == ".pdf"
    assert result["experiences_found"] >= 1
    assert result["skills_found"] >= 3


def test_docx_import(tmp_path: Path, sample_text: str, monkeypatch: pytest.MonkeyPatch):
    from src.common import resume_import as ri

    monkeypatch.setattr(ri, "CANDIDATE_DIR", tmp_path / "candidate")
    monkeypatch.setattr(ri, "CONFIG_DIR", tmp_path / "config")
    monkeypatch.setattr(ri, "RESUMES_MASTER", tmp_path / "resumes")
    (tmp_path / "candidate").mkdir()
    (tmp_path / "config").mkdir()
    (tmp_path / "resumes").mkdir()
    (tmp_path / "config" / "candidate_profile.yaml").write_text(
        "full_name: Jaideep Ponnam\nlocation:\n  city: Austin\n  state: Texas\n",
        encoding="utf-8",
    )
    docx_path = tmp_path / "resume.docx"
    _write_docx(docx_path, sample_text)
    result = import_resume(docx_path, dry_run=True)
    assert result["format"] == ".docx"
    assert result["experiences_found"] >= 1
    assert result["education_blocks_found"] >= 1


def test_date_preservation():
    start, end = parse_dates("Nov 2021 - Dec 2023")
    assert start == "Nov 2021"
    assert end == "Dec 2023"
    start2, end2 = parse_dates("Jul 2025")
    assert start2 == "Jul 2025"
    assert end2 is None
    season, season_end = parse_dates("Spring 2025")
    assert season == "Spring 2025"
    assert season_end is None


def test_role_location_split_preserves_titles():
    from src.common.resume_import import _split_role_location

    assert _split_role_location(
        "Software Engineer (Java & Backend Systems) Atlanta, Georgia, USA"
    ) == ("Software Engineer (Java & Backend Systems)", "Atlanta, Georgia, USA")
    assert _split_role_location("Software Engineer Karnataka, India") == (
        "Software Engineer",
        "Karnataka, India",
    )
    assert _split_role_location("Software Development Intern Hyderabad, India") == (
        "Software Development Intern",
        "Hyderabad, India",
    )


def test_metric_preservation():
    assert extract_metric("Improved runtime by 35% for batch jobs") == "35%"
    assert extract_metric("Built APIs with FastAPI") is None


def test_unsupported_claim_detection():
    from src.common.resume_import import _split_bullets

    resp, meas = _split_bullets(
        [
            "Built Python APIs",
            "Invented world's first quantum LLM chip",
            "Reduced defects by 20%",
        ],
        "sample.md",
        "experience",
    )
    texts = {i["text"]: i for i in resp + meas}
    assert texts["Invented world's first quantum LLM chip"]["claim_type"] == "unsupported"
    assert texts["Reduced defects by 20%"]["metric"] == "20%"


def test_automatic_resume_discovery(tmp_path: Path):
    resume = tmp_path / "My Resume.pdf"
    resume.write_bytes(b"%PDF-1.4")
    # discovery uses suffix only for listing; extract not called
    files = list_resume_files(tmp_path)
    assert files == [resume]
    assert discover_resume(directory=tmp_path) == resume

    (tmp_path / "second.md").write_text("# x", encoding="utf-8")
    with pytest.raises(ValueError, match="Multiple resumes"):
        discover_resume(directory=tmp_path)


def test_profile_validation_reports_issues(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    from src.common import profile_validation as pv
    from src.common import profile_loader as pl
    from src.common import io as io_mod

    candidate = tmp_path / "candidate"
    config = tmp_path / "config"
    candidate.mkdir()
    config.mkdir()
    (config / "candidate_profile.yaml").write_text(
        "full_name: Jaideep Ponnam\nemail: a@b.com\nphone: '123'\n"
        "resume_imported: true\nprofile_reviewed: false\nprofile_approved: false\n"
        "location:\n  city: Austin\n  state: Texas\n"
        "education:\n  - degree: MS CS\n    institution: University of Memphis\n",
        encoding="utf-8",
    )
    (candidate / "verified_experience.yaml").write_text(
        "experiences:\n"
        "  - id: exp_001\n"
        "    employer: A\n"
        "    role: Eng\n"
        "    start_date: Jan 2024\n"
        "    end_date: Dec 2023\n"
        "    responsibilities: []\n"
        "    measurable_results: []\n"
        "    source_file: x.md\n"
        "    source_section: experience\n"
        "    supporting_text: A\n",
        encoding="utf-8",
    )
    (candidate / "skills_inventory.yaml").write_text(
        "skills:\n  - name: Python\n    source_file: x.md\n    source_section: skills\n"
        "    supporting_text: Python\n"
        "  - name: python\n    source_file: x.md\n    source_section: skills\n"
        "    supporting_text: python\n",
        encoding="utf-8",
    )
    (candidate / "project_inventory.yaml").write_text("projects: []\n", encoding="utf-8")
    (candidate / "achievements.yaml").write_text("achievements: []\n", encoding="utf-8")
    (candidate / "certifications.yaml").write_text("certifications: []\n", encoding="utf-8")

    monkeypatch.setattr(pv, "CANDIDATE_DIR", candidate)
    monkeypatch.setattr(pl, "CANDIDATE_DIR", candidate)
    monkeypatch.setattr(pl, "load_candidate_profile", lambda: io_mod.load_yaml(config / "candidate_profile.yaml"))
    monkeypatch.setattr(pv, "load_candidate_profile", lambda: io_mod.load_yaml(config / "candidate_profile.yaml"))

    report = validate_profile()
    severities = {i["severity"] for i in report["issues"]}
    assert "conflicting_dates" in severities
    assert "duplicate_skills" in severities
    assert "empty_experience_descriptions" in severities
    assert report["readiness_percentage"] < 100


def test_tailor_requires_profile_approval(tmp_path: Path):
    profile = CandidateProfile(
        full_name="Jaideep Ponnam",
        education=[EducationRecord(degree="MS CS", institution="University of Memphis")],
        resume_imported=True,
        profile_reviewed=False,
        profile_approved=False,
        skills=[SkillRecord(name="Python")],
        experiences=[
            VerifiedExperience(
                id="exp_001",
                employer="Fixture Labs",
                role="Intern",
                responsibilities=[ClaimItem(text="Built Python APIs")],
            )
        ],
    )
    job = parse_job_dict({"title": "AI Engineer", "company": "Acme", "description": "Python LLM"})
    report = MatchScorer(profile).score(job)
    report.match_score = 80
    with pytest.raises(ValueError, match="reviewed"):
        ResumeTailor(profile, output_dir=tmp_path).tailor(job, report)


def test_normalize_preserves_section_breaks():
    text = normalize_resume_text("May 2025\nMaster's, Computer Science\nimprovin\ng data")
    assert "May 2025" in text
    assert "Master's" in text
    assert "improving data" in text


def test_split_sections_finds_work_experience(sample_text: str):
    sections = split_sections(sample_text)
    assert "experience" in sections
    assert "skills" in sections
    assert "education" in sections
